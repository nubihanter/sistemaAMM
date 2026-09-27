# sales/views.py
import calendar
from datetime import date

import pandas as pd
from django.contrib.auth.decorators import login_required
from django.http import JsonResponse
from django.shortcuts import redirect, render
from django.utils import timezone
from django.utils.http import url_has_allowed_host_and_scheme

from .dashboard_services import gerar_analise_clientes, gerar_metricas_e_graficos
from .models import LogSincronizacao, NotaFiscal, Vendedor


def _obter_redirect_seguro(request, fallback: str) -> str:
    """Retorna uma URL de redirecionamento segura (mesmo host) ou o fallback."""
    candidato = (
        request.POST.get("next")
        or request.GET.get("next")
        or request.META.get("HTTP_REFERER")
        or ""
    ).strip()
    if candidato and url_has_allowed_host_and_scheme(
        url=candidato,
        allowed_hosts={request.get_host()},
        require_https=request.is_secure(),
    ):
        return candidato
    return fallback


def _carregar_df_notas_janela_comercial(ano_selecionado: int, mes_selecionado: int) -> pd.DataFrame:
    """
    Carrega apenas as Notas Fiscais necessárias para o Dashboard Comercial do mês/ano selecionado:
    desde 01/01 do ano anterior (para comparativos YoY, MoM e gráficos de evolução de 12 meses)
    até o último dia do mês selecionado.
    """
    inicio_janela = date(ano_selecionado - 1, 1, 1)
    ultimo_dia_mes = calendar.monthrange(ano_selecionado, mes_selecionado)[1]
    fim_janela = date(ano_selecionado, mes_selecionado, ultimo_dia_mes)

    qs = (
        NotaFiscal.objects.exclude(status="CANCELADA")
        .filter(data_emissao__gte=inicio_janela, data_emissao__lte=fim_janela)
        .values(
            "numero_nota",
            "cliente_nome",
            "cliente_documento",
            "data_emissao",
            "valor_total",
            "vendedor_nome",
        )
    )
    return pd.DataFrame(list(qs))


@login_required
def dashboard_vendas(request):
    user = request.user
    if getattr(user, "role", None) in ("ALMOXARIFADO", "COMPRAS") and not user.is_superuser:
        return redirect("dashboard_estoque")
    hoje = timezone.now().date()

    try:
        mes_selecionado = int(request.GET.get("mes", hoje.month))
        if not 1 <= mes_selecionado <= 12:
            mes_selecionado = hoje.month
    except ValueError:
        mes_selecionado = hoje.month
    try:
        ano_selecionado = int(request.GET.get("ano", hoje.year))
    except ValueError:
        ano_selecionado = hoje.year

    # 1. Carrega as notas faturadas dentro da janela de análise do período selecionado
    df = _carregar_df_notas_janela_comercial(ano_selecionado, mes_selecionado)

    vendedores_ativos = {
        str(v).strip().upper()
        for v in Vendedor.objects.filter(ativo=True).values_list("nome_hardness", flat=True)
        if v
    }

    # 2. Monta a lista de vendedores disponíveis usando consulta distinta indexada no banco
    nomes_na_base = {
        str(v).strip().upper()
        for v in NotaFiscal.objects.exclude(status="CANCELADA")
        .values_list("vendedor_nome", flat=True)
        .distinct()
        if v and str(v).strip().upper() not in ("", "NAN", "NONE")
    }
    vendedores_disponiveis = sorted(v for v in nomes_na_base if v in vendedores_ativos)

    # 3. REGRA DE PERFIL
    if user.is_vendedor:
        vendedora_selecionada = (user.nome_vendedor_erp or user.first_name or user.username).strip().upper()
        pode_selecionar = False
    else:
        vendedora_selecionada = request.GET.get("visao", "EMPRESA").strip().upper()
        pode_selecionar = True

    pode_ver_margem = bool(getattr(user, "is_admin", False))

    # 4. Gera métricas e gráficos Plotly
    (
        metricas,
        graf_evolucao,
        graf_qtd,
        graf_ranking,
        graf_historico_metas,
        graf_top_clientes,
        tabela_notas,
    ) = gerar_metricas_e_graficos(
        df=df,
        vendedora_selecionada=vendedora_selecionada,
        mes_selecionado=mes_selecionado,
        ano_selecionado=ano_selecionado,
        incluir_margem_admin=pode_ver_margem,
        gerar_graficos=True,
    )

    context = {
        "vendedora_selecionada": vendedora_selecionada,
        "vendedores_disponiveis": vendedores_disponiveis,
        "pode_selecionar": pode_selecionar,
        "pode_ver_margem": pode_ver_margem,
        "mes_selecionado": mes_selecionado,
        "ano_selecionado": ano_selecionado,
        "metricas": metricas,
        "grafico_evolucao": graf_evolucao,
        "grafico_qtd": graf_qtd,
        "grafico_ranking": graf_ranking,
        "grafico_historico_metas": graf_historico_metas,
        "grafico_top_clientes": graf_top_clientes,
        "tabela_notas": tabela_notas,
        "meses": [
            (1, "Jan"),
            (2, "Fev"),
            (3, "Mar"),
            (4, "Abr"),
            (5, "Mai"),
            (6, "Jun"),
            (7, "Jul"),
            (8, "Ago"),
            (9, "Set"),
            (10, "Out"),
            (11, "Nov"),
            (12, "Dez"),
        ],
        "anos": [hoje.year, hoje.year - 1, hoje.year - 2],
    }
    return render(request, "sales/dashboard.html", context)


@login_required
def analise_clientes_view(request):
    user = request.user
    if getattr(user, "role", None) in ("ALMOXARIFADO", "COMPRAS") and not user.is_superuser:
        return redirect("dashboard_estoque")

    # 1. Regra de Perfil: Vendedor só pode visualizar sua própria carteira atual
    if user.is_vendedor:
        visao_selecionada = (user.nome_vendedor_erp or user.first_name or user.username).strip().upper()
        pode_selecionar = False
    else:
        visao_selecionada = request.GET.get("visao", "EMPRESA").strip().upper()
        pode_selecionar = True

    # 2. Carrega todas as notas válidas do banco (incluindo cliente_documento para consolidação por CNPJ)
    notas_qs = NotaFiscal.objects.exclude(status="CANCELADA").values(
        "cliente_nome",
        "cliente_documento",
        "vendedor_nome",
        "data_emissao",
        "valor_total",
    )
    df = pd.DataFrame(list(notas_qs))

    # 3. Monta a lista de vendedores ativos para o seletor
    vendedores_qs = (
        Vendedor.objects.filter(ativo=True)
        .exclude(nome_hardness__in=["", "NAN", "NONE", "DESCONHECIDO"])
        .values_list("nome_hardness", flat=True)
        .order_by("nome_hardness")
    )

    vendedores_disponiveis = list(vendedores_qs)

    if not vendedores_disponiveis and not df.empty and "vendedor_nome" in df.columns:
        nomes_unicos = [
            str(v).strip().upper()
            for v in df["vendedor_nome"].dropna().unique()
            if str(v).strip().upper() not in ["", "NAN", "NONE", "DESCONHECIDO"]
        ]
        vendedores_disponiveis = sorted(nomes_unicos)

    # 4. Captura filtros opcionais de Curva e Status
    curvas_selecionadas = request.GET.getlist("curva") or ["AA", "A", "B"]
    status_selecionados = request.GET.getlist("status") or ["Novo", "Ativo", "Em Risco", "Inativo"]

    # 5. Executa os cálculos e gera os gráficos
    kpis, graf_status, graf_matriz, df_tabela = gerar_analise_clientes(
        df=df,
        vendedora_selecionada=visao_selecionada,
        filtros_curva=curvas_selecionadas,
        filtros_status=status_selecionados,
        gerar_graficos=True,
    )

    tabela_linhas = df_tabela.to_dict(orient="records") if not df_tabela.empty else []

    context = {
        "vendedores_disponiveis": vendedores_disponiveis,
        "visao_selecionada": visao_selecionada,
        "pode_selecionar": pode_selecionar,
        "kpis": kpis,
        "grafico_status": graf_status,
        "grafico_matriz": graf_matriz,
        "tabela_clientes": tabela_linhas,
        "curvas_disponiveis": ["AA", "A", "B", "C", "D"],
        "status_disponiveis": ["Novo", "Ativo", "Em Risco", "Inativo"],
        "curvas_selecionadas": curvas_selecionadas,
        "status_selecionados": status_selecionados,
    }
    return render(request, "sales/analise_clientes.html", context)


@login_required
def exportar_vendas_excel_view(request):
    """Exporta os KPIs e Notas Fiscais do Dashboard Comercial em Excel (.xlsx)."""
    from io import BytesIO
    from django.http import HttpResponse
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    user = request.user
    if getattr(user, "role", None) in ("ALMOXARIFADO", "COMPRAS") and not user.is_superuser:
        return redirect("dashboard_estoque")

    hoje = timezone.now().date()
    try:
        mes_selecionado = int(request.GET.get("mes", hoje.month))
        if not 1 <= mes_selecionado <= 12:
            mes_selecionado = hoje.month
    except ValueError:
        mes_selecionado = hoje.month
    try:
        ano_selecionado = int(request.GET.get("ano", hoje.year))
    except ValueError:
        ano_selecionado = hoje.year

    if user.is_vendedor:
        vendedora_selecionada = (user.nome_vendedor_erp or user.first_name or user.username).strip().upper()
    else:
        vendedora_selecionada = request.GET.get("visao", "EMPRESA").strip().upper()

    pode_ver_margem = bool(getattr(user, "is_admin", False))

    df = _carregar_df_notas_janela_comercial(ano_selecionado, mes_selecionado)

    metricas, _, _, _, _, _, tabela_notas = gerar_metricas_e_graficos(
        df=df,
        vendedora_selecionada=vendedora_selecionada,
        mes_selecionado=mes_selecionado,
        ano_selecionado=ano_selecionado,
        incluir_margem_admin=pode_ver_margem,
        gerar_graficos=False,
    )

    wb = Workbook()
    ws_kpis = wb.active
    ws_kpis.title = "Resumo Comercial"

    header_fill = PatternFill(start_color="0F172A", end_color="0F172A", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True, size=11)
    bold_font = Font(bold=True, size=10)
    thin_border = Border(
        left=Side(style="thin", color="CBD5E1"),
        right=Side(style="thin", color="CBD5E1"),
        top=Side(style="thin", color="CBD5E1"),
        bottom=Side(style="thin", color="CBD5E1"),
    )

    ws_kpis.append(["Indicador Comercial", "Valor"])
    for col_idx in (1, 2):
        cell = ws_kpis.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="left" if col_idx == 1 else "right", vertical="center")

    linhas_resumo = [
        ("Visão / Escopo", vendedora_selecionada),
        ("Período (Mês/Ano)", f"{mes_selecionado:02d}/{ano_selecionado}"),
        ("Faturamento Realizado", metricas.get("total_vendas", "R$ 0,00")),
        ("Meta Mensal", metricas.get("meta_total", "R$ 0,00")),
        ("Atingimento da Meta", metricas.get("percentual_meta", "0.0%")),
        ("Status do Ritmo", metricas.get("status_meta", "-")),
        ("Projeção Linear de Fechamento", metricas.get("projecao_fechamento", "R$ 0,00")),
        ("Projeção vs Meta", metricas.get("projecao_pct_meta", "0.0%")),
        ("Comparativo Mês Anterior (MoM)", metricas.get("mom_delta_str", "-")),
        ("Comparativo Mesmo Mês Ano Anterior (YoY)", metricas.get("yoy_delta_str", "-")),
        ("Notas Emitidas no Mês", int(metricas.get("num_vendas", 0))),
        ("Clientes Únicos Atendidos", int(metricas.get("num_clientes", 0))),
        ("Ticket Médio por Nota", metricas.get("ticket_medio", "R$ 0,00")),
    ]
    if pode_ver_margem and metricas.get("tem_margem"):
        linhas_resumo.extend(
            [
                ("Custo Total CMV dos Itens", metricas.get("custo_total_mes", "R$ 0,00")),
                ("Lucro Bruto Estimado", metricas.get("margem_bruta_valor", "R$ 0,00")),
                ("Margem Bruta (%)", metricas.get("margem_bruta_pct", "0.0%")),
            ]
        )

    for label, val in linhas_resumo:
        ws_kpis.append([label, val])
        r = ws_kpis.max_row
        ws_kpis.cell(row=r, column=1).font = bold_font
        ws_kpis.cell(row=r, column=1).border = thin_border
        c_val = ws_kpis.cell(row=r, column=2)
        c_val.border = thin_border
        c_val.alignment = Alignment(horizontal="right", vertical="center")

    ws_kpis.column_dimensions["A"].width = 42
    ws_kpis.column_dimensions["B"].width = 28

    # Aba 2: Notas Fiscais do Período
    ws_nfs = wb.create_sheet(title="Notas Fiscais do Mês")
    colunas_nf = ["Data Emissão", "Número NF", "Cliente (Consolidado por CNPJ)", "Vendedor", "Valor Total (R$)"]
    ws_nfs.append(colunas_nf)
    for col_idx in range(1, len(colunas_nf) + 1):
        c = ws_nfs.cell(row=1, column=col_idx)
        c.fill = header_fill
        c.font = header_font
        c.alignment = Alignment(horizontal="center", vertical="center")

    for nf in tabela_notas:
        ws_nfs.append(
            [
                nf.get("data_str", ""),
                nf.get("numero_nota", ""),
                nf.get("cliente_nome", ""),
                nf.get("vendedor_nome", ""),
                float(nf.get("valor_total", 0)),
            ]
        )
        r = ws_nfs.max_row
        for c_idx in range(1, 6):
            ws_nfs.cell(row=r, column=c_idx).border = thin_border
        ws_nfs.cell(row=r, column=5).number_format = "R$ #,##0.00"

    larguras = [16, 16, 46, 22, 20]
    for idx, larg in enumerate(larguras, start=1):
        ws_nfs.column_dimensions[get_column_letter(idx)].width = larg

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    nome_arq = f"relatorio_comercial_{vendedora_selecionada.lower()}_{mes_selecionado:02d}_{ano_selecionado}.xlsx"
    response = HttpResponse(
        buffer.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = f'attachment; filename="{nome_arq}"'
    return response


@login_required
def exportar_clientes_excel_view(request):
    """Exporta a matriz de inteligência de carteira de clientes em Excel (.xlsx)."""
    from io import BytesIO
    from django.http import HttpResponse
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    user = request.user
    if getattr(user, "role", None) in ("ALMOXARIFADO", "COMPRAS") and not user.is_superuser:
        return redirect("dashboard_estoque")

    if user.is_vendedor:
        visao_selecionada = (user.nome_vendedor_erp or user.first_name or user.username).strip().upper()
    else:
        visao_selecionada = request.GET.get("visao", "EMPRESA").strip().upper()

    curvas_selecionadas = request.GET.getlist("curva") or ["AA", "A", "B"]
    status_selecionados = request.GET.getlist("status") or ["Novo", "Ativo", "Em Risco", "Inativo"]

    notas_qs = NotaFiscal.objects.exclude(status="CANCELADA").values(
        "cliente_nome",
        "cliente_documento",
        "vendedor_nome",
        "data_emissao",
        "valor_total",
    )
    df = pd.DataFrame(list(notas_qs))

    kpis, _, _, df_tabela = gerar_analise_clientes(
        df=df,
        vendedora_selecionada=visao_selecionada,
        filtros_curva=curvas_selecionadas,
        filtros_status=status_selecionados,
        gerar_graficos=False,
    )

    wb = Workbook()
    ws = wb.active
    ws.title = "Carteira de Clientes"

    header_fill = PatternFill(start_color="0F172A", end_color="0F172A", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True, size=11)
    thin_border = Border(
        left=Side(style="thin", color="CBD5E1"),
        right=Side(style="thin", color="CBD5E1"),
        top=Side(style="thin", color="CBD5E1"),
        bottom=Side(style="thin", color="CBD5E1"),
    )

    colunas = [
        "Empresa (Consolidada por CNPJ)",
        "Vendedor Atual (Últ. Compra)",
        "Curva ABC",
        "Status Saúde",
        "Fat. Médio Mensal 6m (R$)",
        "Dias Inativos",
        "Última Venda",
        "Fat. Histórico Total (R$)",
        "Qtd NFs",
    ]
    ws.append(colunas)
    for col_idx in range(1, len(colunas) + 1):
        c = ws.cell(row=1, column=col_idx)
        c.fill = header_fill
        c.font = header_font
        c.alignment = Alignment(horizontal="center", vertical="center")

    if not df_tabela.empty:
        for row in df_tabela.to_dict(orient="records"):
            ws.append(
                [
                    row.get("cliente_nome", ""),
                    row.get("Vendedora", "-"),
                    row.get("Curva", ""),
                    row.get("Status", ""),
                    float(row.get("Ticket_Medio_6m", 0)),
                    int(row.get("Dias_Inatividade", 0)),
                    row.get("Ultima_Venda_str", ""),
                    float(row.get("Faturamento_Total", 0)),
                    int(row.get("Num_Vendas", 0)),
                ]
            )
            r = ws.max_row
            for c_idx in range(1, len(colunas) + 1):
                ws.cell(row=r, column=c_idx).border = thin_border
            ws.cell(row=r, column=5).number_format = "R$ #,##0.00"
            ws.cell(row=r, column=8).number_format = "R$ #,##0.00"

    larguras = [46, 24, 12, 16, 24, 16, 16, 24, 14]
    for idx, larg in enumerate(larguras, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = larg

    # Aba 2: Resumo KPIs da Carteira
    ws_resumo = wb.create_sheet(title="Resumo da Carteira")
    ws_resumo.append(["Indicador", "Valor"])
    for col_idx in (1, 2):
        c = ws_resumo.cell(row=1, column=col_idx)
        c.fill = header_fill
        c.font = header_font

    for titulo, val in [
        ("Escopo / Carteira", visao_selecionada),
        ("Curvas Selecionadas", ", ".join(curvas_selecionadas)),
        ("Status Selecionados", ", ".join(status_selecionados)),
        ("Total de Clientes na Carteira", kpis.get("total_clientes", 0)),
        ("Clientes Novos no Ano", kpis.get("clientes_novos", 0)),
        ("Clientes Ativos (≤30d)", kpis.get("clientes_ativos", 0)),
        ("Clientes Em Risco (31-90d)", kpis.get("clientes_risco", 0)),
        ("Potencial Mensal Em Risco", kpis.get("potencial_risco_mensal", "R$ 0,00")),
        ("Clientes Inativos (>90d)", kpis.get("clientes_inativos", 0)),
        ("Clientes Alta Prioridade (AA/A/B)", kpis.get("prioritarios", 0)),
        ("Faturamento Acumulado da Carteira", kpis.get("faturamento_carteira", "R$ 0,00")),
    ]:
        ws_resumo.append([titulo, val])

    ws_resumo.column_dimensions["A"].width = 36
    ws_resumo.column_dimensions["B"].width = 28

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    nome_arq = f"carteira_clientes_{visao_selecionada.lower()}_{timezone.now().strftime('%Y%m%d')}.xlsx"
    response = HttpResponse(
        buffer.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = f'attachment; filename="{nome_arq}"'
    return response


@login_required
def painel_sincronizacao_view(request):
    """
    Painel de controle para disparar sincronizações manuais (Hardness, PipeRun, CA EPI)
    e auditar logs de execução tanto do painel quanto do agendador (APScheduler).
    """
    from django.db.models import Max
    from django_apscheduler.models import DjangoJob, DjangoJobExecution
    from inventory.models import CertificadoAprovacao, ProdutoEPI
    from .models import MetaVendedor

    user = request.user
    if not getattr(user, "is_admin", False):
        if getattr(user, "role", None) in ("ALMOXARIFADO", "COMPRAS"):
            return redirect("dashboard_estoque")
        return redirect("dashboard_vendas")

    filtro_tipo = request.GET.get("tipo", "").strip()
    filtro_status = request.GET.get("status", "").strip()

    logs_qs = LogSincronizacao.objects.all()
    if filtro_tipo:
        logs_qs = logs_qs.filter(tipo=filtro_tipo)
    if filtro_status:
        logs_qs = logs_qs.filter(status=filtro_status)

    logs_sincronizacao = list(logs_qs[:50])
    tem_sync_em_andamento = any(log.status == "EM_ANDAMENTO" for log in logs_sincronizacao)

    # Última execução de cada tipo
    ultimos_por_tipo = {}
    for codigo_tipo, _ in LogSincronizacao.TIPO_CHOICES:
        ultimos_por_tipo[codigo_tipo] = (
            LogSincronizacao.objects.filter(tipo=codigo_tipo).order_by("-iniciado_em").first()
        )

    # Jobs agendados no django_apscheduler
    try:
        jobs_agendados = list(DjangoJob.objects.all().order_by("id"))
        execucoes_scheduler = list(
            DjangoJobExecution.objects.select_related("job").order_by("-run_time")[:25]
        )
    except Exception:
        jobs_agendados = []
        execucoes_scheduler = []

    # Resumo do estado atual do banco
    hoje = timezone.now().date()
    resumo_banco = {
        "total_notas": NotaFiscal.objects.exclude(status="CANCELADA").count(),
        "ultima_nota_emissao": NotaFiscal.objects.aggregate(Max("data_emissao"))["data_emissao__max"],
        "ultima_nota_sync": NotaFiscal.objects.aggregate(Max("data_sincronizacao"))["data_sincronizacao__max"],
        "total_produtos": ProdutoEPI.objects.count(),
        "ultimo_estoque_sync": ProdutoEPI.objects.aggregate(Max("data_atualizacao"))["data_atualizacao__max"],
        "metas_mes_atual": MetaVendedor.objects.filter(ano=hoje.year, mes=hoje.month).count(),
        "ultima_meta_sync": MetaVendedor.objects.aggregate(Max("data_atualizacao"))["data_atualizacao__max"],
        "total_cas": CertificadoAprovacao.objects.count(),
        "cas_com_validade": CertificadoAprovacao.objects.filter(data_validade__isnull=False).count(),
        "cas_sem_validade": CertificadoAprovacao.objects.filter(data_validade__isnull=True).count(),
    }

    context = {
        "logs_sincronizacao": logs_sincronizacao,
        "tem_sync_em_andamento": tem_sync_em_andamento,
        "ultimos_por_tipo": ultimos_por_tipo,
        "jobs_agendados": jobs_agendados,
        "execucoes_scheduler": execucoes_scheduler,
        "resumo_banco": resumo_banco,
        "filtro_tipo": filtro_tipo,
        "filtro_status": filtro_status,
        "tipos_choices": LogSincronizacao.TIPO_CHOICES,
        "status_choices": LogSincronizacao.STATUS_CHOICES,
    }
    return render(request, "sales/painel_sincronizacao.html", context)


@login_required
def status_sincronizacao_json_view(request):
    """Retorna o status atual das sincronizações em andamento para polling assíncrono do painel."""
    user = request.user
    if not getattr(user, "is_admin", False):
        return JsonResponse({"erro": "Sem permissão."}, status=403)

    em_andamento_qs = LogSincronizacao.objects.filter(status="EM_ANDAMENTO").order_by("-iniciado_em")
    logs_ativos = [
        {
            "id": log.id,
            "tipo": log.tipo,
            "tipo_display": log.get_tipo_display(),
            "iniciado_em": log.iniciado_em.strftime("%H:%M:%S"),
        }
        for log in em_andamento_qs[:10]
    ]
    ultimo = LogSincronizacao.objects.order_by("-iniciado_em").first()
    ultimo_dados = None
    if ultimo:
        ultimo_dados = {
            "id": ultimo.id,
            "tipo": ultimo.tipo,
            "tipo_display": ultimo.get_tipo_display(),
            "status": ultimo.status,
            "mensagem": ultimo.mensagem,
            "duracao_segundos": ultimo.duracao_segundos,
            "registros_criados": ultimo.registros_criados,
            "registros_atualizados": ultimo.registros_atualizados,
        }

    return JsonResponse(
        {
            "em_andamento": len(logs_ativos) > 0,
            "logs_em_andamento": logs_ativos,
            "ultimo_log": ultimo_dados,
        }
    )


@login_required
def disparar_sincronizacao_view(request):
    """
    Dispara manualmente uma tarefa de sincronização solicitada no painel.
    Por padrão executa em background (thread assíncrona) para evitar timeout HTTP,
    com proteção contra disparos duplicados simultâneos.
    """
    from django.contrib import messages
    from inventory.services import sincronizar_estoque_e_itens_rapido, sincronizar_vencimentos_ca
    from .services import (
        disparar_sincronizacao_background,
        registrar_execucao_sincronizacao,
        sincronizar_desde_ultimo_registro,
        sincronizar_metas_piperun,
        sincronizar_notas_hardness,
    )

    if request.method != "POST":
        return redirect("painel_sincronizacao")

    user = request.user
    if not getattr(user, "is_admin", False):
        messages.error(
            request,
            "Apenas o perfil Administrador possui permissão para acessar ou disparar sincronizações.",
        )
        if getattr(user, "role", None) in ("ALMOXARIFADO", "COMPRAS"):
            return redirect("dashboard_estoque")
        return redirect("dashboard_vendas")

    acao = request.POST.get("acao", "").strip()
    modo_sincrono = request.POST.get("modo", "").strip().lower() == "sincrono"
    nome_usuario = user.username

    tipo_sync = None
    funcao_sync = None

    if acao == "notas_hardness":
        tipo_sync = "NOTAS_HARDNESS"
        data_inicio = request.POST.get("data_inicio", "").strip()
        data_fim = request.POST.get("data_fim", "").strip()
        if data_inicio:
            if "-" in data_inicio and len(data_inicio) == 10:
                partes = data_inicio.split("-")
                data_inicio = f"{partes[2]}/{partes[1]}/{partes[0]}"
            if data_fim and "-" in data_fim and len(data_fim) == 10:
                partes_f = data_fim.split("-")
                data_fim = f"{partes_f[2]}/{partes_f[1]}/{partes_f[0]}"
            else:
                data_fim = timezone.now().strftime("%d/%m/%Y")

            funcao_sync = lambda di=data_inicio, df_str=data_fim: sincronizar_notas_hardness(
                data_inicio=di, data_fim=df_str
            )
        else:
            funcao_sync = sincronizar_desde_ultimo_registro

    elif acao == "estoque_hardness":
        tipo_sync = "ESTOQUE_HARDNESS"
        try:
            dias = max(5, min(365, int(request.POST.get("dias_retroativos", 60))))
        except ValueError:
            dias = 60
        funcao_sync = lambda d=dias: sincronizar_estoque_e_itens_rapido(dias_retroativos_padrao=d)

    elif acao == "metas_piperun":
        tipo_sync = "METAS_PIPERUN"
        funcao_sync = lambda: sincronizar_metas_piperun(forcar_api=True)

    elif acao == "ca_epi":
        tipo_sync = "CONSULTA_CA"
        forcar = request.POST.get("forcar_todos") in ("1", "true", "on", "True")

        def _exec_ca(f=forcar):
            r = sincronizar_vencimentos_ca(forcar_todos=f, max_workers=6)
            return {
                "criados": 0,
                "atualizados": r.get("atualizados", 0),
                "mensagem": (
                    f"Processados: {r.get('total', 0)} | Atualizados: {r.get('atualizados', 0)} | "
                    f"Vencidos: {r.get('vencidos', 0)} | Não encontrados: {r.get('nao_encontrados', 0)}"
                ),
            }

        funcao_sync = _exec_ca

    elif acao == "completa":
        tipo_sync = "COMPLETA"

        def _exec_completa():
            nf_c, nf_a = sincronizar_desde_ultimo_registro()
            res_est = sincronizar_estoque_e_itens_rapido(dias_retroativos_padrao=60)
            mt_c, mt_a = sincronizar_metas_piperun(forcar_api=True)
            tot_c = nf_c + res_est.get("estoque_criados", 0) + res_est.get("itens_criados", 0) + mt_c
            tot_a = nf_a + res_est.get("estoque_atualizados", 0) + res_est.get("itens_atualizados", 0) + mt_a
            return {
                "criados": tot_c,
                "atualizados": tot_a,
                "mensagem": (
                    f"NFs: +{nf_c}/{nf_a} | Estoque: +{res_est.get('estoque_criados', 0)}/{res_est.get('estoque_atualizados', 0)} | "
                    f"Itens: +{res_est.get('itens_criados', 0)}/{res_est.get('itens_atualizados', 0)} | Metas: +{mt_c}/{mt_a}"
                ),
            }

        funcao_sync = _exec_completa

    else:
        messages.warning(request, "Ação de sincronização não reconhecida.")
        return redirect(_obter_redirect_seguro(request, "/sales/sincronizacao/"))

    try:
        if modo_sincrono:
            log, _ = registrar_execucao_sincronizacao(
                tipo=tipo_sync,
                funcao_sync=funcao_sync,
                origem="MANUAL_PAINEL",
                usuario=nome_usuario,
            )
            messages.success(
                request,
                f"Sincronização ({log.get_tipo_display()}) concluída com sucesso! {log.mensagem}",
            )
        else:
            log, iniciou = disparar_sincronizacao_background(
                tipo=tipo_sync,
                funcao_sync=funcao_sync,
                origem="MANUAL_PAINEL",
                usuario=nome_usuario,
            )
            if iniciou:
                messages.info(
                    request,
                    f"Sincronização ({log.get_tipo_display()}) iniciada em segundo plano! "
                    f"O painel atualizará automaticamente assim que a execução terminar.",
                )
            else:
                messages.warning(
                    request,
                    f"Já existe uma sincronização ({log.get_tipo_display()}) em andamento "
                    f"iniciada às {log.iniciado_em.strftime('%H:%M:%S')}. Aguarde a conclusão.",
                )
    except Exception as exc:
        messages.error(request, f"Falha durante a sincronização ({acao}): {exc}")

    return redirect(_obter_redirect_seguro(request, "/sales/sincronizacao/"))