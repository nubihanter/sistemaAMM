from datetime import timedelta
from decimal import Decimal
from io import BytesIO

from django.contrib.auth import get_user_model
from django.test import TestCase
from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook

from inventory.models import Categoria, ItemVenda, ProdutoEPI
from sales.models import ContaPagar

User = get_user_model()


class EstoqueFornecedoresEParadosTests(TestCase):
    def setUp(self):
        self.hoje = timezone.now().date()
        self.admin_user = User.objects.create_user(
            username="admin_est",
            password="123",
            role=User.Role.ADMINISTRADOR,
        )
        cat = Categoria.objects.create(nome="EPIS GERAL")
        self.prod_asc = ProdutoEPI.objects.create(
            sku="EPI-ASC-1",
            nome="LUVA NITRILICA",
            marca="MARCA CRESCENTE",
            categoria=cat,
            estoque_atual=50,
            preco_custo=Decimal("10.00"),
            preco_venda=Decimal("25.00"),
        )
        self.prod_queda = ProdutoEPI.objects.create(
            sku="EPI-QUEDA-1",
            nome="OCULOS PROTECAO",
            marca="MARCA CAINDO",
            categoria=cat,
            estoque_atual=40,
            preco_custo=Decimal("8.00"),
            preco_venda=Decimal("20.00"),
        )
        self.prod_parado = ProdutoEPI.objects.create(
            sku="EPI-PARADO-99",
            nome="BOTINA PARADA",
            marca="MARCA PARADA",
            categoria=cat,
            estoque_atual=120,
            preco_custo=Decimal("55.00"),
            preco_venda=Decimal("110.00"),
        )

        # Vendas de MARCA CRESCENTE subindo nos últimos 3 meses (10 -> 30 -> 80)
        ItemVenda.objects.create(
            id_item_erp="IT-1",
            numero_nota="NF-1",
            data_emissao=self.hoje - timedelta(days=75),
            codigo_produto="EPI-ASC-1",
            descricao_produto="LUVA NITRILICA",
            marca="MARCA CRESCENTE",
            quantidade=Decimal("10"),
            valor_unitario=Decimal("25.00"),
            valor_total=Decimal("250.00"),
        )
        ItemVenda.objects.create(
            id_item_erp="IT-2",
            numero_nota="NF-2",
            data_emissao=self.hoje - timedelta(days=45),
            codigo_produto="EPI-ASC-1",
            descricao_produto="LUVA NITRILICA",
            marca="MARCA CRESCENTE",
            quantidade=Decimal("30"),
            valor_unitario=Decimal("25.00"),
            valor_total=Decimal("750.00"),
        )
        ItemVenda.objects.create(
            id_item_erp="IT-3",
            numero_nota="NF-3",
            data_emissao=self.hoje - timedelta(days=10),
            codigo_produto="EPI-ASC-1",
            descricao_produto="LUVA NITRILICA",
            marca="MARCA CRESCENTE",
            quantidade=Decimal("80"),
            valor_unitario=Decimal("25.00"),
            valor_total=Decimal("2000.00"),
        )

        # Vendas de MARCA CAINDO caindo nos últimos 3 meses (90 -> 40 -> 5)
        ItemVenda.objects.create(
            id_item_erp="IT-4",
            numero_nota="NF-4",
            data_emissao=self.hoje - timedelta(days=75),
            codigo_produto="EPI-QUEDA-1",
            descricao_produto="OCULOS PROTECAO",
            marca="MARCA CAINDO",
            quantidade=Decimal("90"),
            valor_unitario=Decimal("20.00"),
            valor_total=Decimal("1800.00"),
        )
        ItemVenda.objects.create(
            id_item_erp="IT-5",
            numero_nota="NF-5",
            data_emissao=self.hoje - timedelta(days=45),
            codigo_produto="EPI-QUEDA-1",
            descricao_produto="OCULOS PROTECAO",
            marca="MARCA CAINDO",
            quantidade=Decimal("40"),
            valor_unitario=Decimal("20.00"),
            valor_total=Decimal("800.00"),
        )
        ItemVenda.objects.create(
            id_item_erp="IT-6",
            numero_nota="NF-6",
            data_emissao=self.hoje - timedelta(days=10),
            codigo_produto="EPI-QUEDA-1",
            descricao_produto="OCULOS PROTECAO",
            marca="MARCA CAINDO",
            quantidade=Decimal("5"),
            valor_unitario=Decimal("20.00"),
            valor_total=Decimal("100.00"),
        )

        # Compras de fornecedores no período (ContaPagar)
        ContaPagar.objects.create(
            empresa="AMM EPIS",
            id_titulo_erp="CP-FORN-1",
            fornecedor_nome="FORNECEDOR LIDER EPI",
            centro_custo="COMPRAS/ESTOQUE",
            grupo_conta="OPERACIONAL - COMPRA MERCADORIA",
            data_emissao=self.hoje - timedelta(days=15),
            data_vencimento=self.hoje + timedelta(days=15),
            valor_total=Decimal("4500.00"),
            valor_saldo=Decimal("4500.00"),
            status="A_VENCER",
        )

    def test_aba_fornecedores_rankings_ascensao_queda_e_consumo_mensal(self):
        self.client.force_login(self.admin_user)
        resp = self.client.get(reverse("dashboard_estoque") + "?dias=90&aba=fornecedores")
        self.assertEqual(resp.status_code, 200)
        self.assertEqual(resp.context["aba_ativa"], "fornecedores")

        comprados = resp.context["tabela_fornecedores_comprados"]
        self.assertEqual(len(comprados), 1)
        self.assertEqual(comprados[0]["Fornecedor"], "FORNECEDOR LIDER EPI")
        self.assertEqual(comprados[0]["Valor_Total_Comprado"], 4500.0)

        vendas_forn = resp.context["tabela_fornecedores_vendas"]
        marcas_map = {f["Fornecedor"]: f for f in vendas_forn}
        self.assertIn("MARCA CRESCENTE", marcas_map)
        self.assertIn("MARCA CAINDO", marcas_map)

        # 120 unidades em 3 meses (90 dias) => consumo médio mensal = 40.0 un/mês
        self.assertAlmostEqual(marcas_map["MARCA CRESCENTE"]["Consumo_Medio_Mensal_Qtd"], 40.0, places=1)
        self.assertEqual(marcas_map["MARCA CRESCENTE"]["Status_Tendencia"], "EM ASCENSÃO")
        self.assertEqual(marcas_map["MARCA CAINDO"]["Status_Tendencia"], "EM QUEDA")

        nomes_asc = [f["Fornecedor"] for f in resp.context["fornecedores_em_ascensao"]]
        nomes_queda = [f["Fornecedor"] for f in resp.context["fornecedores_em_queda"]]
        self.assertIn("MARCA CRESCENTE", nomes_asc)
        self.assertIn("MARCA CAINDO", nomes_queda)

    def test_exportar_parados_excel_excludes_cost_and_tied_capital(self):
        self.client.force_login(self.admin_user)
        resp = self.client.get(reverse("exportar_parados_excel") + "?dias=90")
        self.assertEqual(resp.status_code, 200)
        self.assertIn("spreadsheetml", resp["Content-Type"])

        wb = load_workbook(BytesIO(resp.content))
        ws = wb.active
        headers = [cell.value for cell in ws[1]]

        # Garante que nenhuma coluna de Custo ou Capital Empatado está presente na exportação
        headers_joined = " | ".join(str(h or "").upper() for h in headers)
        self.assertNotIn("CUSTO", headers_joined)
        self.assertNotIn("CAPITAL", headers_joined)
        self.assertNotIn("EMPATADO", headers_joined)

        skus_exportados = [ws.cell(row=r, column=1).value for r in range(2, ws.max_row + 1)]
        self.assertIn("EPI-PARADO-99", skus_exportados)
