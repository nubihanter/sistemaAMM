import csv
from datetime import datetime
from django.contrib import messages
from django.contrib.auth.decorators import login_required
from django.http import HttpResponse, JsonResponse
from django.shortcuts import render, redirect
from django.views.decorators.http import require_POST

from .models import Categoria, ProdutoEPI
from .dashboard_services import gerar_analise_estoque_e_compras
from .services import sincronizar_estoque_e_itens_rapido


@login_required
def dashboard_estoque_view(request):
    try:
        dias_analise = int(request.GET.get("dias", 90))
    except ValueError:
        dias_analise = 90

    try:
        cobertura_meses = float(request.GET.get("cobertura", 2.0))
    except ValueError:
        cobertura_meses = 2.0

    try:
        teto_meses = float(request.GET.get("teto", 4.0))
    except ValueError:
        teto_meses = 4.0

    cliente_selecionado = request.GET.get("cliente", "").strip()
    marca_selecionada = request.GET.get("marca", "TODAS").strip()
    status_filtro = request.GET.get("status", "TODOS").strip()
    apenas_criticos = request.GET.get("criticos", "") == "1"
    aba_ativa = request.GET.get("aba", "compras").strip()

    pode_ver_valores_financeiros = bool(getattr(request.user, "is_admin", False))

    dados = gerar_analise_estoque_e_compras(
        dias_analise=dias_analise,
        cobertura_meses=cobertura_meses,
        teto_meses=teto_meses,
        cliente_selecionado=cliente_selecionado,
        marca_selecionada=marca_selecionada,
        status_filtro=status_filtro,
        apenas_criticos=apenas_criticos,
    )

    if not pode_ver_valores_financeiros and dados.get("tabela_ranking"):
        dados["tabela_ranking"] = sorted(dados["tabela_ranking"], key=lambda x: x.get("Posicao_Qtd", 9999))

    context = {
        **dados,
        "pode_ver_valores_financeiros": pode_ver_valores_financeiros,
        "dias_analise": dias_analise,
        "cobertura_meses": int(cobertura_meses) if cobertura_meses.is_integer() else cobertura_meses,
        "teto_meses": int(teto_meses) if teto_meses.is_integer() else teto_meses,
        "marca_selecionada": marca_selecionada,
        "status_filtro": status_filtro,
        "apenas_criticos": apenas_criticos,
        "aba_ativa": aba_ativa,
        "opcoes_dias": [
            (30, "Últimos 30 dias"),
            (60, "Últimos 60 dias"),
            (90, "Últimos 90 dias (Padrão)"),
            (180, "Últimos 6 meses (180d)"),
            (365, "Último Ano (365d)"),
        ],
        "opcoes_cobertura": [1, 2, 3, 4, 6],
        "opcoes_teto": [3, 4, 6, 9, 12],
        "opcoes_status": [
            ("TODOS", "Todos os Status"),
            ("CRÍTICO / RUPTURA", "🔴 Crítico / Ruptura (Urgente)"),
            ("COMPRAR", "🟠 Comprar (Abaixo da Meta)"),
            ("ESPORÁDICO", "🟣 Item Esporádico (Avaliar)"),
            ("ACIMA DO NECESSÁRIO", "🔵 Acima do Necessário (Excesso)"),
            ("PARADO (SEM GIRO)", "⚫ Parado (Sem Giro no Período)"),
            ("SAUDÁVEL", "🟢 Estoque Saudável"),
        ],
    }
    return render(request, "inventory/dashboard_estoque.html", context)


@login_required
@require_POST
def atualizar_parametros_produto_view(request):
    """
    Permite estabelecer o Estoque Mínimo, Estoque Máximo e marcar/desmarcar
    o produto como Crítico (que não pode faltar) diretamente pelo dashboard.
    """
    sku = request.POST.get("sku", "").strip()
    nome_fallback = request.POST.get("nome", sku).strip()
    marca_fallback = request.POST.get("marca", "N/A").strip()

    if not sku:
        return JsonResponse({"ok": False, "error": "SKU não informado."}, status=400)

    cat_geral, _ = Categoria.objects.get_or_create(nome="GERAL")
    produto, _ = ProdutoEPI.objects.get_or_create(
        sku=sku,
        defaults={
            "nome": nome_fallback or sku,
            "marca": marca_fallback,
            "categoria": cat_geral,
        },
    )

    acao = request.POST.get("acao", "salvar_modal")
    if acao == "toggle_critico":
        produto.item_critico = not produto.item_critico
        if produto.item_critico and produto.estoque_minimo == 0:
            try:
                sugestao_min = int(float(request.POST.get("sugestao_minimo", 0)))
            except ValueError:
                sugestao_min = 0
            if sugestao_min > 0:
                produto.estoque_minimo = sugestao_min
        produto.save(update_fields=["item_critico", "estoque_minimo", "data_atualizacao"])
    else:
        try:
            est_min = max(0, int(request.POST.get("estoque_minimo", 0)))
        except ValueError:
            est_min = produto.estoque_minimo

        try:
            est_max = max(0, int(request.POST.get("estoque_maximo", 0)))
        except ValueError:
            est_max = produto.estoque_maximo

        item_critico = request.POST.get("item_critico") in ("on", "true", "1", "True")

        produto.estoque_minimo = est_min
        produto.estoque_maximo = est_max
        produto.item_critico = item_critico
        produto.save(update_fields=["estoque_minimo", "estoque_maximo", "item_critico", "data_atualizacao"])

    if request.headers.get("x-requested-with") == "XMLHttpRequest":
        return JsonResponse(
            {
                "ok": True,
                "sku": produto.sku,
                "item_critico": produto.item_critico,
                "estoque_minimo": produto.estoque_minimo,
                "estoque_maximo": produto.estoque_maximo,
            }
        )

    messages.success(
        request,
        f"Parâmetros do produto {produto.sku} atualizados (Crítico: {'Sim' if produto.item_critico else 'Não'} | Mín: {produto.estoque_minimo} | Máx: {produto.estoque_maximo}).",
    )
    next_url = request.POST.get("next") or request.META.get("HTTP_REFERER") or "/inventory/dashboard/"
    return redirect(next_url)


@login_required
@require_POST
def sincronizar_estoque_manual_view(request):
    """Dispara a sincronização incremental de estoque e itens vendidos do Hardness."""
    try:
        res = sincronizar_estoque_e_itens_rapido(dias_retroativos_padrao=60)
        messages.success(
            request,
            f"Sincronização concluída! Estoque: {res['estoque_criados']} novos / {res['estoque_atualizados']} atualizados | "
            f"Itens de Venda: {res['itens_criados']} novos / {res['itens_atualizados']} atualizados.",
        )
    except Exception as e:
        messages.error(request, f"Erro ao sincronizar com o Hardness: {e}")

    next_url = request.POST.get("next") or request.META.get("HTTP_REFERER") or "/inventory/dashboard/"
    return redirect(next_url)


@login_required
def exportar_compras_csv_view(request):
    """Exporta a tabela de projeção de compras e diagnóstico de estoque em Excel (.xlsx)."""
    from io import BytesIO
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    pode_ver_valores_financeiros = bool(getattr(request.user, "is_admin", False))

    try:
        dias_analise = int(request.GET.get("dias", 90))
    except ValueError:
        dias_analise = 90
    try:
        cobertura_meses = float(request.GET.get("cobertura", 2.0))
    except ValueError:
        cobertura_meses = 2.0
    try:
        teto_meses = float(request.GET.get("teto", 4.0))
    except ValueError:
        teto_meses = 4.0

    marca_selecionada = request.GET.get("marca", "TODAS").strip()
    status_filtro = request.GET.get("status", "TODOS").strip()
    apenas_criticos = request.GET.get("criticos", "") == "1"

    dados = gerar_analise_estoque_e_compras(
        dias_analise=dias_analise,
        cobertura_meses=cobertura_meses,
        teto_meses=teto_meses,
        marca_selecionada=marca_selecionada,
        status_filtro=status_filtro,
        apenas_criticos=apenas_criticos,
    )

    wb = Workbook()
    ws = wb.active
    ws.title = "Plano de Compras"

    cabecalho = [
        "Código (SKU)",
        "Descrição do Produto",
        "Marca",
        "Unidade",
        "CA",
        "Item Crítico",
        "Curva ABC",
        "Padrão Demanda",
        f"Qtd Vendida ({dias_analise}d)",
        "Pedidos no Período",
        "Clientes Distintos",
        "Média Mensal (un)",
        "Tendência (un/mês)",
        "Estoque Atual",
        "Em Compra (OC)",
        "Estoque Físico",
        "Estoque Mínimo",
        f"Meta Linear ({cobertura_meses:g}M)",
        f"Teto Excesso ({teto_meses:g}M)",
        "Cobertura Atual (Meses)",
        "Status Estoque",
        "Sugestão Compra (un)",
    ]
    if pode_ver_valores_financeiros:
        cabecalho.extend(["Custo Unit. Ref (R$)", "Valor Estimado Compra (R$)"])
    cabecalho.extend(["Qtd Excesso (un)", "Última Venda"])

    ws.append(cabecalho)

    header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True, size=10)
    thin_border = Border(
        left=Side(style="thin", color="E2E8F0"),
        right=Side(style="thin", color="E2E8F0"),
        top=Side(style="thin", color="E2E8F0"),
        bottom=Side(style="thin", color="E2E8F0"),
    )

    for col_idx in range(1, len(cabecalho) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    status_fills = {
        "CRÍTICO / RUPTURA": PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid"),
        "COMPRAR": PatternFill(start_color="FFEDD5", end_color="FFEDD5", fill_type="solid"),
        "ESPORÁDICO": PatternFill(start_color="EDE9FE", end_color="EDE9FE", fill_type="solid"),
        "ACIMA DO NECESSÁRIO": PatternFill(start_color="E0F2FE", end_color="E0F2FE", fill_type="solid"),
        "PARADO (SEM GIRO)": PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid"),
        "SAUDÁVEL": PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid"),
    }

    for row_idx, r in enumerate(dados["tabela_compras"], start=2):
        linha = [
            str(r["codigo_produto"]),
            r["nome"],
            r["marca"],
            r["unidade_medida"],
            r["ca__numero_ca"],
            "SIM" if r["item_critico"] else ("SUGERIDO" if r["Sugestao_Critico"] else "NÃO"),
            r["Curva_ABC"],
            r["Padrao_Demanda"],
            int(round(float(r["Qtd_Vendida"]))),
            int(r["Num_Pedidos"]),
            int(r["Num_Clientes"]),
            round(float(r["Vendas_Media_Mes"]), 1),
            round(float(r["Tendencia_Mes"]), 1),
            int(r["estoque_atual"]),
            int(r["qtd_ordem_compra"]),
            int(r["estoque_fisico"]),
            int(r["estoque_minimo"]),
            int(r["Meta_Estoque"]),
            int(r["Teto_Estoque"]),
            round(float(r["Cobertura_Meses"]), 1),
            r["Status"],
            int(r["Qtd_Sugerida_Compra"]),
        ]
        if pode_ver_valores_financeiros:
            linha.extend(
                [
                    round(float(r["Custo_Ref"]), 2),
                    round(float(r["Valor_Compra_Estimado"]), 2),
                ]
            )
        linha.extend([int(r["Qtd_Excesso"]), r["Ultima_Venda_Str"]])
        ws.append(linha)

        # Formatação numérica e visual da linha
        ws.cell(row=row_idx, column=12).number_format = "0.0"
        ws.cell(row=row_idx, column=13).number_format = "+0.0;-0.0;0.0"
        ws.cell(row=row_idx, column=20).number_format = "0.0"
        status_cell = ws.cell(row=row_idx, column=21)
        if r["Status"] in status_fills:
            status_cell.fill = status_fills[r["Status"]]
            status_cell.font = Font(bold=True, size=9)

        if int(r["Qtd_Sugerida_Compra"]) > 0:
            ws.cell(row=row_idx, column=22).font = Font(bold=True, color="DC2626")

        if pode_ver_valores_financeiros:
            ws.cell(row=row_idx, column=23).number_format = '"R$" #,##0.00'
            ws.cell(row=row_idx, column=24).number_format = '"R$" #,##0.00'

        for c_i in range(1, len(cabecalho) + 1):
            ws.cell(row=row_idx, column=c_i).border = thin_border

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    ws.row_dimensions[1].height = 26

    # Ajuste automático de largura das colunas
    for col_idx in range(1, len(cabecalho) + 1):
        col_letter = get_column_letter(col_idx)
        header_len = len(str(cabecalho[col_idx - 1]))
        max_len = header_len
        for row_idx in range(2, min(ws.max_row + 1, 200)):
            val = ws.cell(row=row_idx, column=col_idx).value
            if val is not None:
                max_len = max(max_len, len(str(val)))
        ws.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 45)

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    filename = f"planejamento_estoque_compras_amm_{datetime.now().strftime('%Y%m%d_%H%M')}.xlsx"
    response = HttpResponse(
        buffer.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    return response

